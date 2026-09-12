from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..core.constantes import (
    COL_A_DECOUVRIR, COL_AJUSTEMENT, COL_CV_ATS, COL_CV_DESIGN, COL_DATE, COL_ENTREPRISE, COL_LETTRE, COL_LIEN,
    COL_LINKEDIN, COL_NOTES, COL_POINTS_FORTS, COL_SCORE, COL_SCORE_INITIAL, COL_STATUT, COL_TITRE, COL_VERDICT,
)

LARGEURS_COLONNES = {
    COL_DATE: 12,
    COL_ENTREPRISE: 28,
    COL_TITRE: 40,
    COL_SCORE: 10,
    COL_POINTS_FORTS: 45,
    COL_A_DECOUVRIR: 45,
    COL_VERDICT: 50,
    COL_LIEN: 45,
    COL_SCORE_INITIAL: 12,
    COL_AJUSTEMENT: 28,
    COL_CV_DESIGN: 30,
    COL_CV_ATS: 30,
    COL_LINKEDIN: 50,
    COL_LETTRE: 60,
    COL_STATUT: 14,
    COL_NOTES: 30,
}
LARGEUR_PAR_DEFAUT = 18

VERT = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
ORANGE = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
ROUGE = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def colorer_score_technique(feuille, df, seuil):
    if df.empty or COL_SCORE not in df.columns:
        return
    colonne = get_column_letter(df.columns.get_loc(COL_SCORE) + 1)
    plage = f"{colonne}2:{colonne}{len(df) + 1}"
    valeur = f'VALUE(LEFT({colonne}2,FIND("/",{colonne}2)-1))'
    feuille.conditional_formatting.add(plage, FormulaRule(formula=[f"{valeur}>={seuil}"], fill=VERT, stopIfTrue=True))
    feuille.conditional_formatting.add(plage, FormulaRule(formula=[f"AND({valeur}>=5,{valeur}<{seuil})"], fill=ORANGE, stopIfTrue=True))
    feuille.conditional_formatting.add(plage, FormulaRule(formula=[f"{valeur}<5"], fill=ROUGE, stopIfTrue=True))


def appliquer_style_suivi(feuille, df, seuil):
    if df.empty:
        return
    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = feuille.dimensions
    for index, colonne in enumerate(df.columns, start=1):
        lettre = get_column_letter(index)
        feuille.column_dimensions[lettre].width = LARGEURS_COLONNES.get(colonne, LARGEUR_PAR_DEFAUT)
        feuille[f"{lettre}1"].font = Font(bold=True)
    colorer_score_technique(feuille, df, seuil)
